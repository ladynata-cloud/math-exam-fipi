#!/usr/bin/env python3
"""Build downloadable Soviet math videos and voiceover scripts from frozen media.

Run after tools/soviet-math-render.cjs and its --check have passed:
  python tools/soviet-math-package.py --out /absolute/output/directory

By default, package all current topics in batches of 30. --first/--last keep
stable course numbers when packaging later additions. Inputs are the course,
narration.json, and media/manifest.json; video bytes are never modified.
Requires python-docx and Node. The output DOCX files still require render/visual
inspection before delivery (see the documents skill in the authoring workflow).
"""
import argparse
import hashlib
import json
import os
import re
import subprocess
import textwrap
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

WORDS_PER_MINUTE = 120


def words(text):
    return len(re.findall(r'[А-Яа-яЁёA-Za-z0-9]+(?:-[А-Яа-яЁёA-Za-z0-9]+)*', text))


def reading_seconds(text):
    return words(text) * 60 / WORDS_PER_MINUTE


def clock(seconds, srt=False):
    ms = round(seconds * 1000)
    h, rest = divmod(ms, 3600000)
    m, rest = divmod(rest, 60000)
    s, ms = divmod(rest, 1000)
    if srt:
        return f'{h:02}:{m:02}:{s:02},{ms:03}'
    fraction = f'{ms:03}'.rstrip('0') or '0'
    prefix = f'{h:02}:' if h else ''
    return f'{prefix}{m:02}:{s:02},{fraction}'


def fmt(seconds):
    return f'{seconds:.1f}'.replace('.', ',')


def digest(path):
    h = hashlib.sha256()
    with path.open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()


def segments(lesson, script):
    video = lesson['video']
    timeline = video['timeline']
    assert len(timeline) == len(script['steps']) == len(lesson['plan']['steps']), script['id']
    assert timeline and video.get('introSeconds', 0) > 0, 'Obsolete video without introduction: ' + script['id']
    assert abs(timeline[0]['questionAt'] - video['introSeconds']) < .001
    out = [dict(step=None, kind='intro', start=0, end=timeline[0]['questionAt'], text=script['intro'].strip())]
    for n, (window, speech) in enumerate(zip(timeline, script['steps'])):
        end = timeline[n + 1]['questionAt'] if n + 1 < len(timeline) else video['finalAt']
        for kind, start, stop in [('question', window['questionAt'], window['explanationAt']), ('explanation', window['explanationAt'], end)]:
            out.append(dict(step=n+1, kind=kind, start=start, end=stop, text=speech[kind].strip()))
    out.append(dict(step=None, kind='final', start=video['finalAt'], end=video['seconds'], text=script['final'].strip()))
    for cue in out:
        assert cue['text'] and cue['end'] > cue['start'], (script['id'], cue)
        assert reading_seconds(cue['text']) <= cue['end']-cue['start']+.001, (script['id'], cue['kind'], 'Speech does not fit')
    for a, b in zip(out, out[1:]):
        assert abs(a['end'] - b['start']) < .001, script['id']
    return out


def subtitle_cues(cues):
    """Keep every speech window aligned with the video; show at most two lines."""
    result = []
    for cue in cues:
        lines = textwrap.wrap(cue['text'], width=42, break_long_words=False, break_on_hyphens=False)
        chunks = ['\n'.join(lines[i:i+2]) for i in range(0, len(lines), 2)]
        at = cue['start']
        for i, chunk in enumerate(chunks):
            end = cue['end'] if i == len(chunks)-1 else at + reading_seconds(chunk)
            assert end > at and end <= cue['end'] + .001
            result.append(dict(start=at, end=end, text=chunk))
            at = end
    return result


def source_lines(lesson):
    source = lesson['video']['source']
    reference = f"{source['authors']}. {source['bookTitle']}. {source['year']}."
    location = f"Раздел: {source['section']}."
    if source.get('pages'):
        location += f" Страницы раздела: {source['pages']}."
    return [reference, location,
            'Авторское условие по теме учебника. Номер упражнения учебника не присваивается.', source['url']]


def add_meta(doc, text, after=3):
    p=doc.add_paragraph(text, 'Metadata')
    p.paragraph_format.space_after=Pt(after)
    return p


def clean_title(text):
    return re.sub(r'[^\w\s]', '', text).strip()


def configure_doc(doc):
    section=doc.sections[0]
    section.page_width=Inches(8.5); section.page_height=Inches(11)
    section.top_margin=Inches(.7); section.bottom_margin=Inches(.65)
    section.left_margin=Inches(.8); section.right_margin=Inches(.8)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        style=doc.styles[name]
        style.font.name='Arial'; style.font.color.rgb=RGBColor(0,0,0)
        style._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial')
    for style in doc.styles:
        for border in style._element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    normal=doc.styles['Normal']; normal.font.size=Pt(11.5)
    normal.paragraph_format.line_spacing=1.08
    normal.paragraph_format.space_after=Pt(4)
    doc.styles['Title'].font.size=Pt(24)
    doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 1'].font.size=Pt(18)
    doc.styles['Heading 1'].paragraph_format.space_after=Pt(8)
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(0)
    meta=doc.styles.add_style('Metadata',1); meta.base_style=normal
    meta.font.size=Pt(9.5); meta.font.color.rgb=RGBColor(70,70,70)
    meta.paragraph_format.line_spacing=1.02; meta.paragraph_format.space_after=Pt(3)
    cue=doc.styles.add_style('Cue',1); cue.base_style=normal
    cue.paragraph_format.space_before=Pt(4); cue.paragraph_format.space_after=Pt(4)
    cue.paragraph_format.keep_together=True
    footer=section.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=footer.add_run(); field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); run._r.addnext(field)
    footer.style=meta
    doc.core_properties.title='Сценарии озвучки Математика по советским учебникам'
    doc.core_properties.subject='Точные временные окна и готовый текст для озвучки видео'
    doc.core_properties.author='Mathexam.space'


def add_listing_table(doc, prepared):
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    widths = [.4, 5.55, .95]
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for cell, text in zip(table.rows[0].cells, ['№', 'Тема', 'Время']):
        cell.text = text
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for lesson, _, _ in prepared:
        for cell, text in zip(table.add_row().cells, [f"{lesson['number']:02}", lesson['topic']['title'], clock(lesson['video']['seconds'])]):
            cell.text = text
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                border = OxmlElement('w:'+edge)
                for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    border.set(qn('w:'+key), value)
                borders.append(border)
            props.append(borders)
            margin = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                element = OxmlElement('w:'+edge)
                element.set(qn('w:w'), '75')
                element.set(qn('w:type'), 'dxa')
                margin.append(element)
            props.append(margin)
            if i == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E9EEF4')
                props.append(shade)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1
                if j != 1:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
                    run.font.bold = i == 0


def cue_metadata(cue):
    read = reading_seconds(cue['text'])
    pause = max(0, cue['end']-cue['start']-read)
    kind = {'intro': 'условие', 'question': 'вопрос', 'explanation': 'объяснение', 'final': 'итог'}[cue['kind']]
    return f"{clock(cue['start'])}–{clock(cue['end'])}  ·  {kind}  ·  речь ≈ {fmt(read)} с, пауза ≈ {fmt(pause)} с"


def write_doc(path, prepared):
    doc = Document()
    configure_doc(doc)
    first, last = prepared[0][0]['number'], prepared[-1][0]['number']
    doc.add_paragraph('Сценарии озвучки', 'Title')
    doc.add_paragraph('Математика по советским учебникам', 'Subtitle')
    doc.add_paragraph(f'Ролики с {first} по {last}. Готовый текст для записи голоса поверх учебных видео. Каждый MP4 записан без звука. Номера и названия файлов совпадают в архивах видео, текстов и субтитров.')
    doc.add_paragraph('Начинайте реплику с указанной отметки. Приблизительное время чтения рассчитано для темпа 120 слов в минуту. Остаток окна оставляйте для паузы, чтобы ученик успел рассмотреть решение. Сначала звучит условие, затем вопросы и объяснения, в конце — итог.')
    doc.add_paragraph('Читайте только текст реплик. Заголовки, таймкоды и сведения об учебнике не озвучиваются. Перед записью проговорите сценарий в своём темпе. Отдельные TXT и SRT находятся в архиве сценариев; SRT можно использовать в видеоредакторе.')
    doc.add_paragraph('Список роликов', 'Heading 1')
    add_listing_table(doc, prepared)
    for lesson, stem, cues in prepared:
        doc.add_page_break()
        doc.add_paragraph(f"{lesson['number']:02} {clean_title(lesson['topic']['title'])}", 'Heading 1')
        add_meta(doc, f"Файл {stem}.mp4  ·  Длительность {clock(lesson['video']['seconds'])}", 6)
        for line in source_lines(lesson):
            add_meta(doc, line)
        for cue in cues:
            paragraph = doc.add_paragraph(style='Cue')
            paragraph.paragraph_format.keep_with_next = cue['kind'] == 'question'
            run = paragraph.add_run(cue_metadata(cue)+'\n')
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(70, 70, 70)
            paragraph.add_run(cue['text'])
    doc.save(path)


def listing_text(prepared):
    lines = ['Математика по советским учебникам', 'Все MP4 записаны без звука.',
             'Номер | Файл | Тема | Длительность', '']
    for lesson, stem, _ in prepared:
        lines.append(f"{lesson['number']:02} | {stem}.mp4 | {lesson['topic']['title']} | {clock(lesson['video']['seconds'])}")
    return '\n'.join(lines)+'\n'


def read_inputs(args):
    repo = args.repo.resolve()
    script_paths = args.scripts or [repo/'soviet-math/narration.json']
    scripts = {}
    for path in script_paths:
        for speech in json.loads(path.read_text()):
            assert speech['id'] not in scripts, 'Duplicate script: '+speech['id']
            scripts[speech['id']] = speech
    node = os.environ.get('CODEX_PRIMARY_RUNTIME_NODE', 'node')
    exporter = """
const root=process.argv[1], path=require('node:path'), crypto=require('node:crypto');
const course=require(path.join(root,'soviet-math/course.js'));
const sources=require(path.join(root,'soviet-math/sources.js'));
const result=course.topics.map((topic,i)=>{const plan=course.make(topic.id,0),source=sources.forTopic(topic.id);return {number:i+1,topic,plan,sourceHash:crypto.createHash('sha256').update(JSON.stringify({topic,plan,source})).digest('hex')};});
process.stdout.write(JSON.stringify(result));
"""
    lessons = json.loads(subprocess.run([node, '-e', exporter, str(repo)], check=True, capture_output=True, text=True).stdout)
    last = args.last or len(lessons)
    assert 1 <= args.first <= last <= len(lessons), 'Invalid course number range'
    lessons = [lesson for lesson in lessons if args.first <= lesson['number'] <= last]
    manifest_path = args.manifest or repo/'soviet-math/media/manifest.json'
    manifest = json.loads(manifest_path.read_text())
    videos = {entry['id']: entry for entry in manifest['lessons']}
    assert len(videos) == len(manifest['lessons']), 'Duplicate video IDs'
    media = args.media or manifest_path.parent
    prepared = []
    for lesson in lessons:
        topic_id = lesson['topic']['id']
        assert topic_id in videos and topic_id in scripts, 'Missing video or script: '+topic_id
        lesson['video'] = video = videos[topic_id]
        speech = scripts[topic_id]
        assert lesson['sourceHash'] == video['sourceHash'], 'Video content is stale: '+topic_id
        speech_hash = hashlib.sha256(json.dumps(speech, ensure_ascii=False, separators=(',', ':')).encode()).hexdigest()
        assert speech_hash == video.get('narrationHash'), 'Video narration is stale: '+topic_id
        assert Path(video['file']).name == video['file'], 'Video file must be a basename'
        filename = media/video['file']
        assert filename.is_file() and digest(filename) == video['sha256'], 'Video bytes changed: '+topic_id
        stem = f"{lesson['number']:02}_{topic_id}"
        prepared.append((lesson, stem, segments(lesson, speech)))
    return prepared, media


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--repo', type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument('--manifest', type=Path)
    parser.add_argument('--scripts', type=Path, nargs='+')
    parser.add_argument('--media', type=Path)
    parser.add_argument('--out', type=Path, required=True)
    parser.add_argument('--first', type=int, default=1)
    parser.add_argument('--last', type=int)
    parser.add_argument('--batch-size', type=int, default=30)
    parser.add_argument('--validate-only', action='store_true')
    parser.add_argument('--scripts-only', action='store_true', help='Build the combined TXT/SRT ZIP without rebuilding DOCX or video ZIPs')
    args = parser.parse_args()
    assert args.batch_size > 0
    prepared, media = read_inputs(args)
    summary = dict(lessons=len(prepared), speech_windows=sum(len(cues) for _, _, cues in prepared),
                   spoken_words=sum(words(cue['text']) for _, _, cues in prepared for cue in cues),
                   seconds=sum(lesson['video']['seconds'] for lesson, _, _ in prepared))
    if args.validate_only:
        print(json.dumps(summary, indent=2))
        return
    out = args.out
    out.mkdir(parents=True, exist_ok=True)
    text_dir, srt_dir = out/'texts', out/'subtitles'
    text_dir.mkdir(exist_ok=True)
    srt_dir.mkdir(exist_ok=True)
    readme = ('Математика по советским учебникам\n\n'
              'Каждый MP4 записан без звука и подготовлен для отдельной озвучки.\n'
              'TXT и SRT содержат авторские сценарии озвучки; текст страниц учебников в них не воспроизводится.\n'
              'TXT содержит готовую речь с временными окнами, оценкой чтения при 120 словах в минуту и паузами.\n'
              'SRT содержит ту же речь короткими субтитрами, согласованными с временными окнами ролика.\n'
              'Номера и английские идентификаторы совпадают у MP4, TXT и SRT.\n'
              'Условия и объяснения авторские. Учебник, авторы, год и страницы указывают тематический источник; '
              'номера упражнений учебника этим примерам не присваиваются.\n')
    (out/'README.txt').write_text(readme)
    outputs = []
    for offset in ([] if args.scripts_only else range(0, len(prepared), args.batch_size)):
        batch = prepared[offset:offset+args.batch_size]
        label = f"{batch[0][0]['number']:02}-{batch[-1][0]['number']:02}"
        docpath = out/f'Scripts_{label}.docx'
        write_doc(docpath, batch)
        videopath = out/f'Videos_{label}.zip'
        with zipfile.ZipFile(videopath, 'w', compression=zipfile.ZIP_STORED) as archive:
            for lesson, stem, _ in batch:
                archive.write(media/lesson['video']['file'], stem+'.mp4')
            archive.writestr(f'List_{label}.txt', listing_text(batch))
            archive.writestr('README.txt', readme)
        outputs.extend([str(videopath), str(docpath)])
    subtitle_count = 0
    for lesson, stem, cues in prepared:
        text = [f"{lesson['number']:02} {lesson['topic']['title']}", f'Файл: {stem}.mp4',
                f"Длительность: {clock(lesson['video']['seconds'])}", *source_lines(lesson), '',
                'Темп чтения примерно 120 слов в минуту. Остаток каждого окна — пауза.',
                'Заголовки, таймкоды и сведения об учебнике не озвучиваются.', '']
        for cue in cues:
            text.extend([cue_metadata(cue), cue['text'], ''])
        (text_dir/f'{stem}.txt').write_text('\n'.join(text))
        subtitles = subtitle_cues(cues)
        subtitle_count += len(subtitles)
        srt = []
        for index, cue in enumerate(subtitles, 1):
            srt.extend([str(index), f"{clock(cue['start'], True)} --> {clock(cue['end'], True)}", cue['text'], ''])
        (srt_dir/f'{stem}.srt').write_text('\n'.join(srt))
    label = f"{prepared[0][0]['number']:02}-{prepared[-1][0]['number']:02}"
    scripts_path = out/f'Scripts_{label}.zip'
    with zipfile.ZipFile(scripts_path, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        for _, stem, _ in prepared:
            archive.write(text_dir/f'{stem}.txt', 'TXT/'+stem+'.txt')
            archive.write(srt_dir/f'{stem}.srt', 'SRT/'+stem+'.srt')
        archive.writestr(f'List_{label}.txt', listing_text(prepared))
        archive.writestr('README.txt', readme)
    outputs.append(str(scripts_path))
    summary.update(subtitle_cues=subtitle_count, outputs=outputs)
    (out/'build-report.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
